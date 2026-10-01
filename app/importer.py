"""
Importa un reporte de SAP (MB51, Plant/SL/BIN, listado de reservas, etc.)
detectando automaticamente las columnas relevantes por su encabezado,
sin depender de que el archivo tenga siempre el mismo formato exacto.
"""
import io
import re
import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

import openpyxl
from sqlalchemy.orm import Session

from .models import Material

COLUMN_PATTERNS = {
    "codigo": [r"^material$", r"^material \d*$", r"^articulo$"],
    "descripcion": [r"^description$", r"^descripcion$", r"^descr$", r"texto breve"],
    "storage_location": [r"^storage location$", r"^storage lo$", r"^sloc$", r"^almacen$"],
    "storage_bin": [r"^storage bin$", r"^storage bi", r"^ubicacion$", r"^bin$"],
    "stock": [r"^unrestricted", r"libre utilizacion"],
    "reservado": [r"reserv"],
    "pendiente": [r"on.?order", r"pendiente", r"^open$", r"abierto"],
    "posting_date": [r"posting date", r"fecha de contab", r"^fecha$"],
    "qty": [r"^qty", r"^quantity", r"^cantidad"],
    "mvt": [r"^mvt", r"movement type", r"clase de mov", r"^mov"],
    "unidad": [r"^base unit$", r"^eun$", r"^unidad"],
    "tipo_material": [r"material type", r"tipo de material"],
    "batch": [r"^batch$", r"^lote$"],
}


def _norm(s) -> str:
    s = "" if s is None else str(s)
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", s.lower()).strip()


def _find_header_row(rows, max_scan=15):
    for i, row in enumerate(rows[:max_scan]):
        cells = [_norm(c) for c in row]
        if "material" in cells:
            return i
    return None


def _map_columns(header_row):
    normed = [_norm(c) for c in header_row]
    mapping = {}
    for field, patterns in COLUMN_PATTERNS.items():
        for idx, cell in enumerate(normed):
            if not cell or idx in mapping.values():
                continue
            if any(re.search(p, cell) for p in patterns):
                mapping[field] = idx
                break
    return mapping


def _to_decimal(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float, Decimal)):
        return Decimal(str(v))
    s = str(v).strip().replace(" ", "")
    neg = s.endswith("-")
    s = s.rstrip("-")
    if re.search(r",\d{1,3}$", s) and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.search(r",\d{1,3}$", s):
        s = s.replace(",", ".")
    else:
        s = s.replace(",", "")
    try:
        d = Decimal(s)
        return -d if neg else d
    except InvalidOperation:
        return None


def _to_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    m = re.match(r"^(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{4})$", s)
    if m:
        a, b, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        mm, dd = (b, a) if a > 12 else (a, b)  # SAP suele exportar MM/DD/AAAA
        try:
            return date(y, mm, dd)
        except ValueError:
            return None
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    return None


def import_workbook(db: Session, file_bytes: bytes) -> dict:
    """Lee un .xlsx, detecta columnas y hace upsert sobre la tabla materiales."""
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()

    hi = _find_header_row(rows)
    if hi is None:
        return {"ok": False, "error": "No encontre una columna 'Material' en el archivo."}

    header = rows[hi]
    col = _map_columns(header)
    data_rows = rows[hi + 1:]

    is_movement_report = "posting_date" in col and "qty" in col
    stats = {"nuevos": 0, "actualizados": 0, "ingresos": 0, "filas": 0, "redeployment": 0, "otros_sloc": 0}
    ultimos_ingresos = {}  # codigo -> (fecha, cantidad)

    for row in data_rows:
        if not row or col.get("codigo") is None:
            continue
        codigo = row[col["codigo"]]
        if codigo is None or str(codigo).strip() == "":
            continue
        codigo = str(codigo).strip()
        stats["filas"] += 1

        if is_movement_report:
            mvt = str(row[col["mvt"]]).strip() if col.get("mvt") is not None else "101"
            if mvt != "101":
                continue
            f = _to_date(row[col["posting_date"]])
            q = _to_decimal(row[col["qty"]])
            if f is None:
                continue
            prev = ultimos_ingresos.get(codigo)
            if prev is None or f > prev[0]:
                ultimos_ingresos[codigo] = (f, abs(q) if q is not None else None)
            continue

        sloc = None
        if col.get("storage_location") is not None and row[col["storage_location"]] is not None:
            sloc = str(row[col["storage_location"]]).strip()

        # Reporte de redeployment (deposito 9001): solo carga ese numero
        # aparte, sin tocar ubicacion/stock de 1117 del mismo material.
        if sloc == "9001":
            material = db.get(Material, codigo)
            is_new = material is None
            if is_new:
                material = Material(codigo=codigo)
                db.add(material)
            if col.get("descripcion") is not None and row[col["descripcion"]] and not material.descripcion:
                material.descripcion = str(row[col["descripcion"]]).strip()
            if col.get("stock") is not None:
                v = _to_decimal(row[col["stock"]])
                if v is not None:
                    material.redeployment_stock = v
            stats["redeployment"] += 1
            stats["nuevos" if is_new else "actualizados"] += 1
            continue

        # Cualquier otro deposito que no sea 1117 (ni 9001): no lo tocamos,
        # para no mezclar ubicaciones/stocks de depositos distintos.
        if sloc and sloc not in ("1117", ""):
            stats["otros_sloc"] += 1
            continue

        material = db.get(Material, codigo)
        is_new = material is None
        if is_new:
            material = Material(codigo=codigo)
            db.add(material)

        if col.get("descripcion") is not None and row[col["descripcion"]]:
            material.descripcion = str(row[col["descripcion"]]).strip()
        if col.get("storage_bin") is not None and row[col["storage_bin"]]:
            material.storage_bin = str(row[col["storage_bin"]]).strip()
        if col.get("stock") is not None:
            v = _to_decimal(row[col["stock"]])
            if v is not None:
                material.stock = v
        if col.get("reservado") is not None:
            v = _to_decimal(row[col["reservado"]])
            if v is not None:
                material.reservado = (material.reservado or 0) + v
        if col.get("pendiente") is not None:
            v = _to_decimal(row[col["pendiente"]])
            if v is not None:
                material.pendiente = (material.pendiente or 0) + v
        if col.get("unidad") is not None and row[col["unidad"]]:
            material.unidad = str(row[col["unidad"]]).strip()
        if col.get("tipo_material") is not None and row[col["tipo_material"]]:
            material.tipo_material = str(row[col["tipo_material"]]).strip()
        if col.get("batch") is not None and row[col["batch"]]:
            material.batch = str(row[col["batch"]]).strip()

        stats["nuevos" if is_new else "actualizados"] += 1

    for codigo, (fecha, cantidad) in ultimos_ingresos.items():
        material = db.get(Material, codigo)
        if material is None:
            continue
        material.ultimo_ingreso_fecha = fecha
        material.ultimo_ingreso_cantidad = cantidad
        stats["ingresos"] += 1

    db.commit()
    stats["ok"] = True
    if is_movement_report:
        stats["tipo_detectado"] = "movimientos (ultimo ingreso)"
    elif stats["redeployment"] and stats["redeployment"] == stats["nuevos"] + stats["actualizados"]:
        stats["tipo_detectado"] = "redeployment (9001)"
    else:
        stats["tipo_detectado"] = "maestro / stock"
    stats["columnas_detectadas"] = list(col.keys())
    return stats
